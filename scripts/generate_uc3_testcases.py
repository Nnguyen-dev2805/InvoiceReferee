"""Generate synthetic UC-03 paper settlement and e-invoice test fixtures."""

from __future__ import annotations

import argparse
import json
import shutil
from copy import deepcopy
from dataclasses import dataclass, field
from datetime import date
from decimal import Decimal
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    KeepTogether,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


COMPANY_NAME = "CÔNG TY TNHH GIẢI PHÁP CÔNG NGHỆ AN THỊNH"
COMPANY_TAX_CODE = "0319999996"
COMPANY_ADDRESS = "128 Nguyễn Văn Linh, Phường Tân Hưng, Thành phố Hồ Chí Minh"
EMPLOYEE_NAME = "NGUYỄN MINH AN"
EMPLOYEE_DEPARTMENT = "Hành chính nhân sự"
ACCOUNTANT_NAME = "TRẦN THỊ THU HÀ"
DIRECTOR_NAME = "LÊ QUANG HUY"
SELLER_SIGNER_NAME = "PHẠM HOÀNG NAM"
SELLER_NAME = "CÔNG TY TNHH THƯƠNG MẠI VÀ DỊCH VỤ MINH PHÁT"
SELLER_TAX_CODE = "0101234565"
SELLER_ADDRESS = "86 Trần Thái Tông, Phường Cầu Giấy, Thành phố Hà Nội"
FONT_DIR = Path(r"C:\Windows\Fonts")

VENDOR_PROFILES = {
    "hoa_don_khach_san": (
        "CÔNG TY TNHH KHÁCH SẠN SÔNG HÀN",
        "0402198455",
        "36 Bạch Đằng, Phường Hải Châu, Thành phố Đà Nẵng",
        "NGUYỄN THỊ THU LAN",
    ),
    "hoa_don_ve_may_bay": (
        "CÔNG TY CỔ PHẦN DỊCH VỤ HÀNG KHÔNG PHƯƠNG NAM",
        "0318125473",
        "18 Hồng Hà, Phường Tân Sơn Hòa, Thành phố Hồ Chí Minh",
        "TRẦN QUỐC BẢO",
    ),
    "hoa_don_taxi": (
        "CÔNG TY CỔ PHẦN VẬN TẢI SAO VIỆT",
        "0109876540",
        "72 Nguyễn Khang, Phường Yên Hòa, Thành phố Hà Nội",
        "LÊ MINH TUẤN",
    ),
    "hoa_don_nha_hang": (
        "CÔNG TY TNHH ẨM THỰC GIA VIÊN",
        "0317462589",
        "145 Nguyễn Đình Chiểu, Phường Bàn Cờ, Thành phố Hồ Chí Minh",
        "VÕ THỊ NGỌC MAI",
    ),
    "hoa_don_phong_hop": (
        "CÔNG TY TNHH TRUNG TÂM HỘI NGHỊ MINH TÂM",
        "0107654326",
        "28 Duy Tân, Phường Cầu Giấy, Thành phố Hà Nội",
        "ĐẶNG HOÀNG LONG",
    ),
    "hoa_don_linh_kien": (
        "CÔNG TY TNHH THIẾT BỊ VĂN PHÒNG PHÚC AN",
        "0316587426",
        "219 Lý Thường Kiệt, Phường Phú Thọ, Thành phố Hồ Chí Minh",
        "PHAN GIA HUY",
    ),
    "hoa_don_nhien_lieu": (
        "CÔNG TY TNHH NHIÊN LIỆU ĐÔNG NAM",
        "0309876547",
        "12 Mai Chí Thọ, Phường An Khánh, Thành phố Hồ Chí Minh",
        "BÙI THANH SƠN",
    ),
    "hoa_don_thue_hoi_truong": (
        "CÔNG TY TNHH SỰ KIỆN THÀNH CÔNG",
        "0104567890",
        "95 Trần Duy Hưng, Phường Yên Hòa, Thành phố Hà Nội",
        "ĐỖ THU HƯƠNG",
    ),
    "hoa_don_in_an": (
        "CÔNG TY TNHH IN VÀ TRUYỀN THÔNG ÁNH DƯƠNG",
        "0315678945",
        "44 Lê Văn Sỹ, Phường Nhiêu Lộc, Thành phố Hồ Chí Minh",
        "NGÔ QUANG VINH",
    ),
    "hoa_don_thiet_bi": (
        "CÔNG TY TNHH GIẢI PHÁP MẠNG HƯNG PHÁT",
        "0106789458",
        "63 Hoàng Quốc Việt, Phường Nghĩa Đô, Thành phố Hà Nội",
        "PHẠM ĐỨC ANH",
    ),
    "hoa_don_van_chuyen": (
        "CÔNG TY TNHH VẬN TẢI THỊNH VƯỢNG",
        "0317894564",
        "27 Đường số 8, Phường Bình Hưng Hòa, Thành phố Hồ Chí Minh",
        "LƯU VĂN HẢI",
    ),
    "hoa_don_van_phong_pham": (
        "CÔNG TY TNHH VĂN PHÒNG PHẨM HÒA BÌNH",
        "0103456789",
        "51 Nguyễn Phong Sắc, Phường Cầu Giấy, Thành phố Hà Nội",
        "HOÀNG THỊ MINH CHÂU",
    ),
    "hoa_don_muc_in": (
        "CÔNG TY TNHH CÔNG NGHỆ MỰC IN VIỆT",
        "0312345673",
        "102 Cộng Hòa, Phường Bảy Hiền, Thành phố Hồ Chí Minh",
        "TRƯƠNG HỮU PHƯỚC",
    ),
    "hoa_don_khong_khai_bao": (
        "CÔNG TY TNHH DỊCH VỤ ẨM THỰC PHỐ XANH",
        "0108765438",
        "16 Phan Chu Trinh, Phường Cửa Nam, Thành phố Hà Nội",
        "ĐINH THỊ HỒNG",
    ),
    "hoa_don_ve_tau": (
        "CÔNG TY CỔ PHẦN DỊCH VỤ ĐƯỜNG SẮT AN PHÚ",
        "0314567897",
        "01 Nguyễn Thông, Phường Xuân Hòa, Thành phố Hồ Chí Minh",
        "NGUYỄN VĂN KHẢI",
    ),
}


@dataclass
class Line:
    description: str
    unit: str
    quantity: Decimal
    unit_price: Decimal

    @property
    def amount(self) -> Decimal:
        return self.quantity * self.unit_price


@dataclass
class TestCase:
    case_id: str
    title: str
    purpose: str
    invoice_number: str
    lines: list[Line]
    tax_rate: Decimal = Decimal("10")
    invoice_date: date = date(2026, 10, 10)
    template_number: str = "1C26TAA"
    serial_number: str = "IR26E"
    seller_name: str = SELLER_NAME
    seller_tax_code: str | None = SELLER_TAX_CODE
    seller_address: str = SELLER_ADDRESS
    seller_signer: str = SELLER_SIGNER_NAME
    buyer_name: str = COMPANY_NAME
    buyer_tax_code: str | None = COMPANY_TAX_CODE
    invoice_subtotal: Decimal | None = None
    invoice_tax: Decimal | None = None
    invoice_total: Decimal | None = None
    invoice_words_value: Decimal | None = None
    report_lines: list[dict[str, Any]] | None = None
    report_total: Decimal | None = None
    report_words_value: Decimal | None = None
    advance_amount: Decimal = Decimal("0")
    declared_difference: Decimal | None = None
    expected_decision: str = "AUTO_PROCESS"
    expected_uncertainty: str = "NONE"
    expected_primary_rule: str | None = None
    expected_findings: list[str] = field(default_factory=list)
    implementation_status: str = "COVERED"
    rationale: str = ""

    def subtotal(self) -> Decimal:
        return self.invoice_subtotal or sum(
            (line.amount for line in self.lines), Decimal("0")
        )

    def tax(self) -> Decimal:
        if self.invoice_tax is not None:
            return self.invoice_tax
        return (self.subtotal() * self.tax_rate / Decimal("100")).quantize(
            Decimal("1")
        )

    def total(self) -> Decimal:
        return self.invoice_total or self.subtotal() + self.tax()

    def settlement_lines(self) -> list[dict[str, Any]]:
        if self.report_lines is not None:
            return self.report_lines
        return [
            {
                "description": self.purpose,
                "invoice_ref": self.invoice_number,
                "amount": self.total(),
                "note": "Đính kèm hóa đơn điện tử",
            }
        ]

    def settlement_total(self) -> Decimal:
        if self.report_total is not None:
            return self.report_total
        return sum(
            (Decimal(str(line["amount"])) for line in self.settlement_lines()),
            Decimal("0"),
        )

    def difference(self) -> Decimal:
        if self.declared_difference is not None:
            return self.declared_difference
        return abs(self.advance_amount - self.settlement_total())


@dataclass
class EvidenceDocument:
    evidence_id: str
    file_stem: str
    document_type: str
    reference: str
    title: str
    description: str
    amount: Decimal = Decimal("0")
    lines: list[Line] = field(default_factory=list)
    tax_rate: Decimal = Decimal("0")
    issuer: str = SELLER_NAME
    seller_tax_code: str = SELLER_TAX_CODE
    issuer_address: str = SELLER_ADDRESS
    issuer_signer: str = SELLER_SIGNER_NAME

    def total(self) -> Decimal:
        if self.document_type == "E_INVOICE":
            subtotal = sum((line.amount for line in self.lines), Decimal("0"))
            tax = (subtotal * self.tax_rate / Decimal("100")).quantize(Decimal("1"))
            return subtotal + tax
        return self.amount


@dataclass
class MultiDocumentCase:
    case_id: str
    title: str
    purpose: str
    report_lines: list[dict[str, Any]]
    evidence_documents: list[EvidenceDocument]
    advance_amount: Decimal
    expected_decision: str = "AUTO_PROCESS"
    expected_uncertainty: str = "NONE"
    expected_primary_rule: str | None = None
    expected_findings: list[str] = field(default_factory=list)
    implementation_status: str = "GAP_EXPECTED"
    rationale: str = ""
    report_total: Decimal | None = None
    declared_difference: Decimal | None = None
    line_expectations: list[dict[str, Any]] = field(default_factory=list)

    def settlement_total(self) -> Decimal:
        if self.report_total is not None:
            return self.report_total
        return sum(
            (Decimal(str(line["amount"])) for line in self.report_lines),
            Decimal("0"),
        )

    def difference(self) -> Decimal:
        if self.declared_difference is not None:
            return self.declared_difference
        return abs(self.advance_amount - self.settlement_total())

    def as_report_case(self) -> TestCase:
        return TestCase(
            case_id=self.case_id,
            title=self.title,
            purpose=self.purpose,
            invoice_number="MULTI",
            lines=[Line("Hồ sơ nhiều chứng từ", "Hồ sơ", Decimal("1"), Decimal("0"))],
            tax_rate=Decimal("0"),
            report_lines=self.report_lines,
            report_total=self.report_total,
            advance_amount=self.advance_amount,
            declared_difference=self.declared_difference,
        )


def money(value: Decimal | int | str) -> str:
    return f"{Decimal(str(value)):,.0f}".replace(",", ".")


def vietnamese_number(value: Decimal | int | str) -> str:
    number = int(Decimal(str(value)))
    if number == 0:
        return "Không đồng"
    digits = [
        "không",
        "một",
        "hai",
        "ba",
        "bốn",
        "năm",
        "sáu",
        "bảy",
        "tám",
        "chín",
    ]

    def triad(n: int, full: bool) -> str:
        hundred = n // 100
        ten = (n % 100) // 10
        unit = n % 10
        parts: list[str] = []
        if hundred or full:
            parts.extend([digits[hundred], "trăm"])
        if ten > 1:
            parts.extend([digits[ten], "mươi"])
            if unit == 1:
                parts.append("mốt")
            elif unit == 4:
                parts.append("tư")
            elif unit == 5:
                parts.append("lăm")
            elif unit:
                parts.append(digits[unit])
        elif ten == 1:
            parts.append("mười")
            if unit == 5:
                parts.append("lăm")
            elif unit:
                parts.append(digits[unit])
        elif unit:
            if hundred or full:
                parts.append("lẻ")
            parts.append(digits[unit])
        return " ".join(parts)

    groups: list[int] = []
    remaining = number
    while remaining:
        groups.append(remaining % 1000)
        remaining //= 1000
    scales = ["", "nghìn", "triệu", "tỷ"]
    words: list[str] = []
    for index in range(len(groups) - 1, -1, -1):
        group = groups[index]
        if group == 0:
            continue
        full = bool(words) and group < 100
        words.append(triad(group, full))
        scale_index = index % 4
        billion_cycle = index // 4
        if scale_index:
            words.append(scales[scale_index])
        if billion_cycle:
            words.extend(["tỷ"] * billion_cycle)
    result = " ".join(words).strip()
    return result[:1].upper() + result[1:] + " đồng"


def _set_repeat_header(row: Any) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def _shade(cell: Any, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _cell_margin(cell: Any, top: int = 90, start: int = 90, bottom: int = 90, end: int = 90) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _cell_text(
    cell: Any,
    text: str,
    *,
    bold: bool = False,
    align: WD_ALIGN_PARAGRAPH = WD_ALIGN_PARAGRAPH.LEFT,
    size: int = 9,
    color: str = "000000",
) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    _cell_margin(cell)


def _configure_doc(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Cm(1.35)
    section.bottom_margin = Cm(1.25)
    section.left_margin = Cm(1.45)
    section.right_margin = Cm(1.45)
    section.header_distance = Cm(0.5)
    section.footer_distance = Cm(0.55)
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(4)
    document.core_properties.title = "Hồ sơ hoàn ứng và chứng từ chi phí"
    document.core_properties.subject = "Hồ sơ kế toán nội bộ"
    document.core_properties.author = "Phòng Tài chính Kế toán"


def _add_labeled_paragraph(document: Document, label: str, value: str, *, bold_value: bool = False) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(2)
    label_run = paragraph.add_run(label)
    label_run.bold = True
    value_run = paragraph.add_run(value)
    value_run.bold = bold_value


def create_report_docx(case: TestCase, output: Path) -> None:
    document = Document()
    _configure_doc(document)

    header = document.add_table(rows=1, cols=2)
    header.alignment = WD_TABLE_ALIGNMENT.CENTER
    _cell_text(header.cell(0, 0), COMPANY_NAME, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=10)
    _cell_text(
        header.cell(0, 1),
        "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM\nĐộc lập - Tự do - Hạnh phúc",
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        size=10,
    )

    date_paragraph = document.add_paragraph(
        f"Thành phố Hồ Chí Minh, ngày {case.invoice_date.day:02d} tháng "
        f"{case.invoice_date.month:02d} năm {case.invoice_date.year}"
    )
    date_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    date_paragraph.paragraph_format.space_after = Pt(4)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(2)
    run = title.add_run("GIẤY ĐỀ NGHỊ HOÀN ỨNG")
    run.bold = True
    run.font.size = Pt(16)
    subtitle = document.add_paragraph(f"Về việc: {case.purpose}")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(6)
    subtitle.runs[0].italic = True

    _add_labeled_paragraph(document, "Kính gửi: ", "Ban Giám đốc " + COMPANY_NAME)
    _add_labeled_paragraph(document, "Người đề nghị: ", EMPLOYEE_NAME)
    _add_labeled_paragraph(document, "Chức vụ: ", "Nhân viên")
    _add_labeled_paragraph(document, "Bộ phận: ", EMPLOYEE_DEPARTMENT)
    _add_labeled_paragraph(document, "Nội dung: ", case.purpose)

    report_lines = case.settlement_lines()
    table = document.add_table(rows=1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    headers = ["STT", "NỘI DUNG", "HÓA ĐƠN", "SỐ TIỀN (VND)", "GHI CHÚ"]
    widths = [Cm(1.0), Cm(6.3), Cm(2.6), Cm(3.3), Cm(3.6)]
    for index, label in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.width = widths[index]
        _shade(cell, "D9EAF2")
        _cell_text(cell, label, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=9)
    _set_repeat_header(table.rows[0])
    for index, line in enumerate(report_lines, start=1):
        cells = table.add_row().cells
        values = [
            str(index),
            str(line["description"]),
            str(line["invoice_ref"]),
            money(line["amount"]),
            str(line.get("note") or ""),
        ]
        for column, value in enumerate(values):
            alignment = (
                WD_ALIGN_PARAGRAPH.RIGHT
                if column == 3
                else WD_ALIGN_PARAGRAPH.CENTER
                if column in {0, 2}
                else WD_ALIGN_PARAGRAPH.LEFT
            )
            _cell_text(cells[column], value, align=alignment, size=9)
    total_cells = table.add_row().cells
    merged = total_cells[0].merge(total_cells[2])
    _cell_text(merged, "Tổng cộng", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=9)
    _cell_text(
        total_cells[3],
        money(case.settlement_total()),
        bold=True,
        align=WD_ALIGN_PARAGRAPH.RIGHT,
        size=9,
    )
    _cell_text(total_cells[4], "", size=9)

    words_value = case.report_words_value or case.settlement_total()
    _add_labeled_paragraph(document, "Bằng chữ: ", vietnamese_number(words_value) + ".")
    _add_labeled_paragraph(
        document,
        "Số tiền đã tạm ứng: ",
        money(case.advance_amount) + " đồng",
        bold_value=True,
    )
    if case.advance_amount > case.settlement_total():
        difference_label = "Số tiền nhân viên hoàn trả công ty: "
    elif case.advance_amount < case.settlement_total():
        difference_label = "Số tiền công ty thanh toán thêm: "
    else:
        difference_label = "Chênh lệch quyết toán: "
    _add_labeled_paragraph(
        document,
        difference_label,
        money(case.difference()) + " đồng",
        bold_value=True,
    )
    _add_labeled_paragraph(
        document,
        "Bằng chữ: ",
        vietnamese_number(case.difference()) + ".",
    )

    signatures = document.add_table(rows=1, cols=3)
    signatures.alignment = WD_TABLE_ALIGNMENT.CENTER
    labels = [
        f"NGƯỜI ĐỀ NGHỊ HOÀN ỨNG\n(Đã ký)\n{EMPLOYEE_NAME}",
        f"KẾ TOÁN\n(Đã ký)\n{ACCOUNTANT_NAME}",
        f"GIÁM ĐỐC\n(Đã ký, đóng dấu)\n{DIRECTOR_NAME}",
    ]
    for index, label in enumerate(labels):
        _cell_text(signatures.cell(0, index), label, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=9)
        signatures.cell(0, index).paragraphs[0].paragraph_format.space_after = Pt(34)

    document.save(output)


def create_invoice_docx(case: TestCase, output: Path) -> None:
    document = Document()
    _configure_doc(document)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("HÓA ĐƠN GIÁ TRỊ GIA TĂNG ĐIỆN TỬ")
    run.bold = True
    run.font.size = Pt(16)
    date_paragraph = document.add_paragraph(
        f"Ngày {case.invoice_date.day:02d} tháng {case.invoice_date.month:02d} "
        f"năm {case.invoice_date.year}"
    )
    date_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    meta = document.add_table(rows=1, cols=3)
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_values = [
        f"Mẫu số: {case.template_number}",
        f"Ký hiệu: {case.serial_number}",
        f"Số: {case.invoice_number}",
    ]
    for index, value in enumerate(meta_values):
        _cell_text(meta.cell(0, index), value, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=9)

    _add_labeled_paragraph(document, "Đơn vị bán hàng: ", case.seller_name, bold_value=True)
    _add_labeled_paragraph(document, "Mã số thuế: ", case.seller_tax_code or "")
    _add_labeled_paragraph(document, "Địa chỉ: ", case.seller_address)
    _add_labeled_paragraph(document, "Đơn vị mua hàng: ", case.buyer_name, bold_value=True)
    _add_labeled_paragraph(document, "Mã số thuế: ", case.buyer_tax_code or "")
    _add_labeled_paragraph(document, "Địa chỉ: ", COMPANY_ADDRESS)
    _add_labeled_paragraph(document, "Hình thức thanh toán: ", "Chuyển khoản")

    table = document.add_table(rows=1, cols=6)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    headers = ["STT", "TÊN HÀNG HÓA, DỊCH VỤ", "ĐVT", "SL", "ĐƠN GIÁ", "THÀNH TIỀN"]
    widths = [Cm(0.9), Cm(6.0), Cm(1.4), Cm(1.1), Cm(3.2), Cm(3.6)]
    for index, label in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.width = widths[index]
        _shade(cell, "D9EAF2")
        _cell_text(cell, label, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=8.5)
    _set_repeat_header(table.rows[0])
    for index, line in enumerate(case.lines, start=1):
        cells = table.add_row().cells
        values = [
            str(index),
            line.description,
            line.unit,
            money(line.quantity),
            money(line.unit_price),
            money(line.amount),
        ]
        for column, value in enumerate(values):
            alignment = (
                WD_ALIGN_PARAGRAPH.RIGHT
                if column in {3, 4, 5}
                else WD_ALIGN_PARAGRAPH.CENTER
                if column in {0, 2}
                else WD_ALIGN_PARAGRAPH.LEFT
            )
            _cell_text(cells[column], value, align=alignment, size=8.5)

    totals = [
        ("Cộng tiền hàng chưa thuế:", case.subtotal(), False),
        (f"Thuế GTGT ({money(case.tax_rate)}%):", case.tax(), False),
        ("TỔNG CỘNG THANH TOÁN:", case.total(), True),
    ]
    for label, value, bold in totals:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.add_run(label).bold = bold
        tab_stops = paragraph.paragraph_format.tab_stops
        tab_stops.add_tab_stop(Cm(16), WD_ALIGN_PARAGRAPH.RIGHT)
        value_run = paragraph.add_run("\t" + money(value) + " VND")
        value_run.bold = bold
    invoice_words = case.invoice_words_value or case.total()
    _add_labeled_paragraph(document, "Số tiền bằng chữ: ", vietnamese_number(invoice_words) + ".")

    signatures = document.add_table(rows=1, cols=2)
    signatures.alignment = WD_TABLE_ALIGNMENT.CENTER
    _cell_text(
        signatures.cell(0, 0),
        "NGƯỜI MUA HÀNG\n(Ký, ghi rõ họ tên)",
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        size=9,
    )
    _cell_text(
        signatures.cell(0, 1),
        f"NGƯỜI BÁN HÀNG\nKý bởi: {case.seller_name}\n"
        f"Người ký: {case.seller_signer}\nKý ngày: {case.invoice_date:%d/%m/%Y}",
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        size=9,
    )
    for cell in signatures.rows[0].cells:
        cell.paragraphs[0].paragraph_format.space_after = Pt(24)
    document.save(output)


def _register_pdf_fonts() -> None:
    pdfmetrics.registerFont(TTFont("TimesVN", str(FONT_DIR / "times.ttf")))
    pdfmetrics.registerFont(TTFont("TimesVN-Bold", str(FONT_DIR / "timesbd.ttf")))
    pdfmetrics.registerFont(TTFont("TimesVN-Italic", str(FONT_DIR / "timesi.ttf")))


def _pdf_styles() -> dict[str, ParagraphStyle]:
    base = getSampleStyleSheet()
    return {
        "body": ParagraphStyle(
            "BodyVN",
            parent=base["BodyText"],
            fontName="TimesVN",
            fontSize=9.5,
            leading=12,
            spaceAfter=3,
        ),
        "small": ParagraphStyle(
            "SmallVN",
            parent=base["BodyText"],
            fontName="TimesVN",
            fontSize=8,
            leading=9.5,
        ),
        "center": ParagraphStyle(
            "CenterVN",
            parent=base["BodyText"],
            fontName="TimesVN",
            fontSize=9.5,
            leading=12,
            alignment=TA_CENTER,
        ),
        "title": ParagraphStyle(
            "TitleVN",
            parent=base["Title"],
            fontName="TimesVN-Bold",
            fontSize=16,
            leading=19,
            alignment=TA_CENTER,
            textColor=colors.black,
            spaceAfter=4,
        ),
        "warning": ParagraphStyle(
            "WarningVN",
            parent=base["BodyText"],
            fontName="TimesVN-Bold",
            fontSize=8,
            leading=10,
            alignment=TA_CENTER,
            textColor=colors.HexColor("#B41E1E"),
            spaceAfter=6,
        ),
        "right": ParagraphStyle(
            "RightVN",
            parent=base["BodyText"],
            fontName="TimesVN",
            fontSize=9,
            alignment=TA_RIGHT,
        ),
    }


def _pdf_table(data: list[list[Any]], widths: list[float], *, header: bool = True) -> Table:
    table = Table(data, colWidths=widths, repeatRows=1 if header else 0, hAlign="CENTER")
    commands = [
        ("GRID", (0, 0), (-1, -1), 0.55, colors.HexColor("#9A9A9A")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]
    if header:
        commands.extend(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#D9EAF2")),
                ("FONTNAME", (0, 0), (-1, 0), "TimesVN-Bold"),
                ("ALIGN", (0, 0), (-1, 0), "CENTER"),
            ]
        )
    table.setStyle(TableStyle(commands))
    return table


def create_report_pdf(case: TestCase, output: Path) -> None:
    styles = _pdf_styles()
    doc = SimpleDocTemplate(
        str(output),
        pagesize=A4,
        leftMargin=15 * mm,
        rightMargin=15 * mm,
        topMargin=12 * mm,
        bottomMargin=14 * mm,
    )
    story: list[Any] = []
    header = Table(
        [
            [
                Paragraph(f"<b>{COMPANY_NAME}</b>", styles["center"]),
                Paragraph(
                    "<b>CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM</b><br/>"
                    "Độc lập - Tự do - Hạnh phúc",
                    styles["center"],
                ),
            ]
        ],
        colWidths=[88 * mm, 88 * mm],
    )
    header.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP")]))
    story.extend(
        [
            header,
            Paragraph(
                f"Thành phố Hồ Chí Minh, ngày {case.invoice_date.day:02d} tháng "
                f"{case.invoice_date.month:02d} năm {case.invoice_date.year}",
                styles["right"],
            ),
            Spacer(1, 2 * mm),
            Paragraph("GIẤY ĐỀ NGHỊ HOÀN ỨNG", styles["title"]),
            Paragraph(f"<i>Về việc: {case.purpose}</i>", styles["center"]),
            Spacer(1, 2 * mm),
            Paragraph(f"<b>Kính gửi:</b> Ban Giám đốc {COMPANY_NAME}", styles["body"]),
            Paragraph(f"<b>Người đề nghị:</b> {EMPLOYEE_NAME}", styles["body"]),
            Paragraph("<b>Chức vụ:</b> Nhân viên", styles["body"]),
            Paragraph(f"<b>Bộ phận:</b> {EMPLOYEE_DEPARTMENT}", styles["body"]),
            Paragraph(f"<b>Nội dung:</b> {case.purpose}", styles["body"]),
            Spacer(1, 2 * mm),
        ]
    )
    rows: list[list[Any]] = [
        [
            Paragraph("STT", styles["small"]),
            Paragraph("NỘI DUNG", styles["small"]),
            Paragraph("HÓA ĐƠN", styles["small"]),
            Paragraph("SỐ TIỀN", styles["small"]),
            Paragraph("GHI CHÚ", styles["small"]),
        ]
    ]
    for index, line in enumerate(case.settlement_lines(), start=1):
        rows.append(
            [
                Paragraph(str(index), styles["small"]),
                Paragraph(str(line["description"]), styles["small"]),
                Paragraph(str(line["invoice_ref"]), styles["small"]),
                Paragraph(money(line["amount"]), styles["small"]),
                Paragraph(str(line.get("note") or ""), styles["small"]),
            ]
        )
    rows.append(
        [
            "",
            Paragraph("<b>Tổng cộng</b>", styles["small"]),
            "",
            Paragraph(f"<b>{money(case.settlement_total())}</b>", styles["small"]),
            "",
        ]
    )
    report_table = _pdf_table(rows, [10 * mm, 61 * mm, 26 * mm, 31 * mm, 48 * mm])
    report_table.setStyle(
        TableStyle(
            [
                ("ALIGN", (0, 1), (0, -1), "CENTER"),
                ("ALIGN", (2, 1), (2, -1), "CENTER"),
                ("ALIGN", (3, 1), (3, -1), "RIGHT"),
                ("SPAN", (0, -1), (0, -1)),
            ]
        )
    )
    story.extend([report_table, Spacer(1, 2 * mm)])
    words_value = case.report_words_value or case.settlement_total()
    story.append(Paragraph(f"<b>Bằng chữ:</b> {vietnamese_number(words_value)}.", styles["body"]))
    story.append(
        Paragraph(
            f"<b>Số tiền đã tạm ứng:</b> {money(case.advance_amount)} đồng",
            styles["body"],
        )
    )
    if case.advance_amount > case.settlement_total():
        label = "Số tiền nhân viên hoàn trả công ty"
    elif case.advance_amount < case.settlement_total():
        label = "Số tiền công ty thanh toán thêm"
    else:
        label = "Chênh lệch quyết toán"
    story.extend(
        [
            Paragraph(f"<b>{label}:</b> {money(case.difference())} đồng", styles["body"]),
            Paragraph(f"<b>Bằng chữ:</b> {vietnamese_number(case.difference())}.", styles["body"]),
            Spacer(1, 6 * mm),
        ]
    )
    signature_data = [
        [
            Paragraph(f"<b>NGƯỜI ĐỀ NGHỊ HOÀN ỨNG</b><br/>(Đã ký)<br/>{EMPLOYEE_NAME}", styles["center"]),
            Paragraph(f"<b>KẾ TOÁN</b><br/>(Đã ký)<br/>{ACCOUNTANT_NAME}", styles["center"]),
            Paragraph(f"<b>GIÁM ĐỐC</b><br/>(Đã ký, đóng dấu)<br/>{DIRECTOR_NAME}", styles["center"]),
        ]
    ]
    signature_table = Table(signature_data, colWidths=[58 * mm] * 3)
    signature_table.setStyle(
        TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("BOTTOMPADDING", (0, 0), (-1, -1), 28)])
    )
    story.append(signature_table)
    doc.build(story)


def create_invoice_pdf(case: TestCase, output: Path) -> None:
    styles = _pdf_styles()
    doc = SimpleDocTemplate(
        str(output),
        pagesize=A4,
        leftMargin=15 * mm,
        rightMargin=15 * mm,
        topMargin=12 * mm,
        bottomMargin=14 * mm,
    )
    story: list[Any] = [
        Paragraph("HÓA ĐƠN GIÁ TRỊ GIA TĂNG ĐIỆN TỬ", styles["title"]),
        Paragraph(
            f"Ngày {case.invoice_date.day:02d} tháng {case.invoice_date.month:02d} "
            f"năm {case.invoice_date.year}",
            styles["center"],
        ),
        Spacer(1, 2 * mm),
    ]
    meta = Table(
        [
            [
                Paragraph(f"<b>Mẫu số:</b> {case.template_number}", styles["small"]),
                Paragraph(f"<b>Ký hiệu:</b> {case.serial_number}", styles["small"]),
                Paragraph(f"<b>Số:</b> {case.invoice_number}", styles["small"]),
            ]
        ],
        colWidths=[58 * mm] * 3,
    )
    story.extend(
        [
            meta,
            Spacer(1, 2 * mm),
            Paragraph(f"<b>Đơn vị bán hàng:</b> {case.seller_name}", styles["body"]),
            Paragraph(f"<b>Mã số thuế:</b> {case.seller_tax_code or ''}", styles["body"]),
            Paragraph(f"<b>Địa chỉ:</b> {case.seller_address}", styles["body"]),
            Spacer(1, 1 * mm),
            Paragraph(f"<b>Đơn vị mua hàng:</b> {case.buyer_name}", styles["body"]),
            Paragraph(f"<b>Mã số thuế:</b> {case.buyer_tax_code or ''}", styles["body"]),
            Paragraph(f"<b>Địa chỉ:</b> {COMPANY_ADDRESS}", styles["body"]),
            Paragraph("<b>Hình thức thanh toán:</b> Chuyển khoản", styles["body"]),
            Spacer(1, 2 * mm),
        ]
    )
    rows: list[list[Any]] = [
        [
            Paragraph("STT", styles["small"]),
            Paragraph("TÊN HÀNG HÓA, DỊCH VỤ", styles["small"]),
            Paragraph("ĐVT", styles["small"]),
            Paragraph("SL", styles["small"]),
            Paragraph("ĐƠN GIÁ", styles["small"]),
            Paragraph("THÀNH TIỀN", styles["small"]),
        ]
    ]
    for index, line in enumerate(case.lines, start=1):
        rows.append(
            [
                Paragraph(str(index), styles["small"]),
                Paragraph(line.description, styles["small"]),
                Paragraph(line.unit, styles["small"]),
                Paragraph(money(line.quantity), styles["small"]),
                Paragraph(money(line.unit_price), styles["small"]),
                Paragraph(money(line.amount), styles["small"]),
            ]
        )
    invoice_table = _pdf_table(rows, [10 * mm, 63 * mm, 15 * mm, 13 * mm, 34 * mm, 41 * mm])
    invoice_table.setStyle(
        TableStyle(
            [
                ("ALIGN", (0, 1), (0, -1), "CENTER"),
                ("ALIGN", (2, 1), (3, -1), "CENTER"),
                ("ALIGN", (4, 1), (5, -1), "RIGHT"),
            ]
        )
    )
    story.extend([invoice_table, Spacer(1, 3 * mm)])
    summary_rows = [
        ["Cộng tiền hàng chưa thuế:", money(case.subtotal())],
        [f"Thuế GTGT ({money(case.tax_rate)}%):", money(case.tax())],
        ["TỔNG CỘNG THANH TOÁN:", money(case.total()) + " VND"],
    ]
    summary = Table(summary_rows, colWidths=[120 * mm, 55 * mm])
    summary.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (-1, -2), "TimesVN"),
                ("FONTNAME", (0, -1), (-1, -1), "TimesVN-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 9.5),
                ("ALIGN", (1, 0), (1, -1), "RIGHT"),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    invoice_words = case.invoice_words_value or case.total()
    story.extend(
        [
            summary,
            Paragraph(f"<b>Số tiền bằng chữ:</b> {vietnamese_number(invoice_words)}.", styles["body"]),
            Spacer(1, 5 * mm),
        ]
    )
    signatures = Table(
        [
            [
                Paragraph("<b>NGƯỜI MUA HÀNG</b><br/>(Ký, ghi rõ họ tên)", styles["center"]),
                Paragraph(
                    f"<b>NGƯỜI BÁN HÀNG</b><br/>{case.seller_name}<br/>"
                    f"Ký bởi: {case.seller_name}<br/>"
                    f"Người ký: {case.seller_signer}<br/>"
                    f"Ký ngày: {case.invoice_date:%d/%m/%Y}",
                    styles["center"],
                ),
            ]
        ],
        colWidths=[87 * mm] * 2,
    )
    signatures.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("BOTTOMPADDING", (0, 0), (-1, -1), 24)]))
    story.append(signatures)
    doc.build(story)


def evidence_as_invoice_case(
    owner: MultiDocumentCase,
    evidence: EvidenceDocument,
) -> TestCase:
    if evidence.document_type != "E_INVOICE" or not evidence.lines:
        raise ValueError(f"{evidence.evidence_id} không phải hóa đơn điện tử")
    return TestCase(
        case_id=f"{owner.case_id}-{evidence.evidence_id}",
        title=evidence.title,
        purpose=evidence.description,
        invoice_number=evidence.reference,
        lines=evidence.lines,
        tax_rate=evidence.tax_rate,
        seller_name=evidence.issuer,
        seller_tax_code=evidence.seller_tax_code,
        seller_address=evidence.issuer_address,
        seller_signer=evidence.issuer_signer,
        advance_amount=Decimal("0"),
    )


def create_alternative_docx(
    owner: MultiDocumentCase,
    evidence: EvidenceDocument,
    output: Path,
) -> None:
    document = Document()
    _configure_doc(document)

    provider = document.add_paragraph(evidence.issuer)
    provider.alignment = WD_ALIGN_PARAGRAPH.CENTER
    provider.runs[0].bold = True
    provider.runs[0].font.size = Pt(11)

    title = document.add_paragraph(evidence.title.upper())
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.runs[0].bold = True
    title.runs[0].font.size = Pt(16)
    reference = document.add_paragraph(f"Số tham chiếu: {evidence.reference}")
    reference.alignment = WD_ALIGN_PARAGRAPH.CENTER

    _add_labeled_paragraph(document, "Ngày chứng từ: ", "10/10/2026")
    _add_labeled_paragraph(document, "Đơn vị/Người nhận: ", COMPANY_NAME, bold_value=True)
    _add_labeled_paragraph(document, "Nội dung: ", evidence.description)
    _add_labeled_paragraph(
        document,
        "Số tiền: ",
        money(evidence.total()) + " VND",
        bold_value=True,
    )
    _add_labeled_paragraph(
        document,
        "Bằng chữ: ",
        vietnamese_number(evidence.total()) + ".",
    )
    payment_method = {
        "BANK_TRANSFER": "Chuyển khoản",
        "E_TICKET": "Thanh toán điện tử",
        "RECEIPT": "Tiền mặt",
        "ACCEPTANCE_MINUTES": "Theo hợp đồng dịch vụ",
    }.get(evidence.document_type, "Theo chứng từ")
    _add_labeled_paragraph(document, "Hình thức thanh toán: ", payment_method)
    if evidence.document_type == "BANK_TRANSFER":
        _add_labeled_paragraph(document, "Trạng thái giao dịch: ", "Thành công", bold_value=True)
        _add_labeled_paragraph(document, "Người chuyển tiền: ", EMPLOYEE_NAME)
        _add_labeled_paragraph(document, "Tài khoản thụ hưởng: ", "0311000123456")
        note_text = "Giao dịch đã được hệ thống ngân hàng ghi nhận hoàn tất."
    elif evidence.document_type == "ACCEPTANCE_MINUTES":
        _add_labeled_paragraph(document, "Kết quả nghiệm thu: ", "Dịch vụ đã hoàn thành và được tiếp nhận", bold_value=True)
        note_text = "Hai bên thống nhất nội dung và giá trị nghiệm thu nêu trên."
    elif evidence.document_type == "E_TICKET":
        _add_labeled_paragraph(document, "Thời gian sử dụng: ", "10/10/2026 08:15")
        _add_labeled_paragraph(document, "Biển kiểm soát: ", "51H-428.36")
        note_text = "Vé được phát hành điện tử và có giá trị cho một lượt sử dụng."
    else:
        _add_labeled_paragraph(document, "Người nộp tiền: ", EMPLOYEE_NAME)
        note_text = "Phiếu thu được lập thành 02 liên có giá trị như nhau."
    note = document.add_paragraph(note_text)
    note.runs[0].italic = True

    signatures = document.add_table(rows=1, cols=2)
    signatures.alignment = WD_TABLE_ALIGNMENT.CENTER
    _cell_text(
        signatures.cell(0, 0),
        f"NGƯỜI LẬP CHỨNG TỪ\n(Đã ký)\n{ACCOUNTANT_NAME}",
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        size=9,
    )
    _cell_text(
        signatures.cell(0, 1),
        f"ĐƠN VỊ XÁC NHẬN\n(Đã ký, đóng dấu)\n{evidence.issuer}",
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        size=9,
    )
    for cell in signatures.rows[0].cells:
        cell.paragraphs[0].paragraph_format.space_after = Pt(50)
    document.save(output)


def create_alternative_pdf(
    owner: MultiDocumentCase,
    evidence: EvidenceDocument,
    output: Path,
) -> None:
    styles = _pdf_styles()
    doc = SimpleDocTemplate(
        str(output),
        pagesize=A4,
        leftMargin=20 * mm,
        rightMargin=20 * mm,
        topMargin=15 * mm,
        bottomMargin=16 * mm,
    )
    story: list[Any] = [
        Paragraph(f"<b>{evidence.issuer}</b>", styles["center"]),
        Spacer(1, 5 * mm),
        Paragraph(evidence.title.upper(), styles["title"]),
        Paragraph(f"Số tham chiếu: {evidence.reference}", styles["center"]),
        Spacer(1, 8 * mm),
        Paragraph("<b>Ngày chứng từ:</b> 10/10/2026", styles["body"]),
        Paragraph(f"<b>Đơn vị/Người nhận:</b> {COMPANY_NAME}", styles["body"]),
        Paragraph(f"<b>Nội dung:</b> {evidence.description}", styles["body"]),
        Paragraph(f"<b>Số tiền:</b> {money(evidence.total())} VND", styles["body"]),
        Paragraph(
            f"<b>Bằng chữ:</b> {vietnamese_number(evidence.total())}.",
            styles["body"],
        ),
        Paragraph(
            "<b>Hình thức thanh toán:</b> "
            + {
                "BANK_TRANSFER": "Chuyển khoản",
                "E_TICKET": "Thanh toán điện tử",
                "RECEIPT": "Tiền mặt",
                "ACCEPTANCE_MINUTES": "Theo hợp đồng dịch vụ",
            }.get(evidence.document_type, "Theo chứng từ"),
            styles["body"],
        ),
    ]
    if evidence.document_type == "BANK_TRANSFER":
        story.extend([
            Paragraph("<b>Trạng thái giao dịch:</b> Thành công", styles["body"]),
            Paragraph(f"<b>Người chuyển tiền:</b> {EMPLOYEE_NAME}", styles["body"]),
            Paragraph("<b>Tài khoản thụ hưởng:</b> 0311000123456", styles["body"]),
            Paragraph("<i>Giao dịch đã được hệ thống ngân hàng ghi nhận hoàn tất.</i>", styles["body"]),
        ])
    elif evidence.document_type == "ACCEPTANCE_MINUTES":
        story.extend([
            Paragraph("<b>Kết quả nghiệm thu:</b> Dịch vụ đã hoàn thành và được tiếp nhận", styles["body"]),
            Paragraph("<i>Hai bên thống nhất nội dung và giá trị nghiệm thu nêu trên.</i>", styles["body"]),
        ])
    elif evidence.document_type == "E_TICKET":
        story.extend([
            Paragraph("<b>Thời gian sử dụng:</b> 10/10/2026 08:15", styles["body"]),
            Paragraph("<b>Biển kiểm soát:</b> 51H-428.36", styles["body"]),
            Paragraph("<i>Vé được phát hành điện tử và có giá trị cho một lượt sử dụng.</i>", styles["body"]),
        ])
    else:
        story.extend([
            Paragraph(f"<b>Người nộp tiền:</b> {EMPLOYEE_NAME}", styles["body"]),
            Paragraph("<i>Phiếu thu được lập thành 02 liên có giá trị như nhau.</i>", styles["body"]),
        ])
    story.append(Spacer(1, 12 * mm))
    signatures = Table(
        [[
            Paragraph(f"<b>NGƯỜI LẬP CHỨNG TỪ</b><br/>(Đã ký)<br/>{ACCOUNTANT_NAME}", styles["center"]),
            Paragraph(f"<b>ĐƠN VỊ XÁC NHẬN</b><br/>(Đã ký, đóng dấu)<br/>{evidence.issuer}", styles["center"]),
        ]],
        colWidths=[80 * mm, 80 * mm],
    )
    signatures.setStyle(
        TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("BOTTOMPADDING", (0, 0), (-1, -1), 50)])
    )
    story.append(signatures)
    doc.build(story)


def base_case(case_id: str, title: str, *, total_kind: str = "meal") -> TestCase:
    if total_kind == "meal":
        lines = [Line("Dịch vụ ăn uống tiếp khách", "Suất", Decimal("10"), Decimal("300000"))]
        purpose = "Chi phí tiếp khách phục vụ dự án Nâng cấp văn phòng 2026"
    elif total_kind == "office":
        lines = [Line("Máy in văn phòng", "Cái", Decimal("1"), Decimal("5000000"))]
        purpose = "Mua thiết bị văn phòng phục vụ hoạt động công ty"
    elif total_kind == "hotel":
        lines = [Line("Phòng nghỉ công tác", "Đêm", Decimal("4"), Decimal("1000000"))]
        purpose = "Chi phí lưu trú trong chuyến công tác"
    else:
        raise ValueError(total_kind)
    return TestCase(
        case_id=case_id,
        title=title,
        purpose=purpose,
        invoice_number=f"0000{case_id[-2:]}01",
        lines=lines,
        tax_rate=Decimal("8"),
        advance_amount=Decimal("4000000"),
    )


def build_cases() -> list[TestCase]:
    cases: list[TestCase] = []

    tc01 = base_case("TC01", "Hợp lệ nhân viên hoàn trả tiền thừa")
    tc01.advance_amount = Decimal("4000000")
    tc01.rationale = "Tổng chi 3.240.000 VND; nhân viên hoàn trả 760.000 VND."
    tc01.expected_findings = ["SET_PAPER_TOTAL_001", "SET_ADV_AMOUNT_001", "SET_SETTLEMENT_DIFF_001", "SET_APPROVAL_001"]
    cases.append(tc01)

    tc02 = base_case("TC02", "Hợp lệ công ty thanh toán thêm", total_kind="office")
    tc02.advance_amount = Decimal("5000000")
    tc02.rationale = "Tổng chi 5.400.000 VND; công ty thanh toán thêm 400.000 VND."
    tc02.expected_findings = ["SET_PAPER_TOTAL_001", "SET_ADV_AMOUNT_001", "SET_SETTLEMENT_DIFF_001", "SET_APPROVAL_001"]
    cases.append(tc02)

    tc03 = base_case("TC03", "Hợp lệ không phát sinh chênh lệch", total_kind="hotel")
    tc03.advance_amount = Decimal("4320000")
    tc03.rationale = "Số tạm ứng bằng đúng tổng chi phí."
    tc03.expected_findings = ["SET_PAPER_TOTAL_001", "SET_ADV_AMOUNT_001", "SET_SETTLEMENT_DIFF_001", "SET_APPROVAL_001"]
    cases.append(tc03)

    tc04 = TestCase(
        case_id="TC04",
        title="Hợp lệ nhiều mặt hàng thiết bị",
        purpose="Mua thiết bị máy tính cho nhân viên",
        invoice_number="00000401",
        lines=[
            Line("Máy tính để bàn", "Bộ", Decimal("5"), Decimal("12000000")),
            Line("Chuột máy tính", "Cái", Decimal("5"), Decimal("200000")),
            Line("Màn hình máy tính", "Cái", Decimal("5"), Decimal("3000000")),
        ],
        tax_rate=Decimal("10"),
        advance_amount=Decimal("85000000"),
        rationale="Tổng 83.600.000 VND; nhân viên hoàn trả 1.400.000 VND.",
        expected_findings=["SET_PAPER_MATCH_001", "SET_PAPER_TOTAL_001", "SET_SETTLEMENT_DIFF_001", "SET_APPROVAL_001"],
    )
    cases.append(tc04)

    tc05 = base_case("TC05", "Tổng các dòng trên báo cáo không khớp")
    tc05.report_lines = [
        {"description": "Chi phí tiếp khách", "invoice_ref": "00000501", "amount": Decimal("2000000"), "note": "Đính kèm hóa đơn"},
        {"description": "Chi phí phòng họp", "invoice_ref": "00000501", "amount": Decimal("1000000"), "note": "Đính kèm hóa đơn"},
    ]
    tc05.report_total = Decimal("3500000")
    tc05.advance_amount = Decimal("4000000")
    tc05.expected_decision = "REQUEST_INFO"
    tc05.expected_uncertainty = "FACTUAL_UNKNOWN"
    tc05.expected_primary_rule = "SET_PAPER_TOTAL_001"
    tc05.expected_findings = ["SET_PAPER_TOTAL_001"]
    tc05.rationale = "Hai dòng cộng 3.000.000 VND nhưng tổng báo cáo ghi 3.500.000 VND."
    cases.append(tc05)

    tc06 = base_case("TC06", "Chênh lệch quyết toán ghi sai")
    tc06.advance_amount = Decimal("4000000")
    tc06.declared_difference = Decimal("700000")
    tc06.expected_decision = "REQUEST_INFO"
    tc06.expected_uncertainty = "FACTUAL_UNKNOWN"
    tc06.expected_primary_rule = "SET_SETTLEMENT_DIFF_001"
    tc06.expected_findings = ["SET_SETTLEMENT_DIFF_001"]
    tc06.rationale = "Chênh lệch đúng là 760.000 VND nhưng báo cáo ghi 700.000 VND."
    cases.append(tc06)

    tc07 = base_case("TC07", "Số tiền bằng chữ trên báo cáo không khớp")
    tc07.report_words_value = Decimal("3500000")
    tc07.expected_decision = "REQUEST_INFO"
    tc07.expected_uncertainty = "FACTUAL_UNKNOWN"
    tc07.expected_primary_rule = "SET_PAPER_CONSISTENCY_001"
    tc07.expected_findings = ["SET_PAPER_CONSISTENCY_001"]
    tc07.rationale = "Tổng bằng số là 3.240.000 VND nhưng bằng chữ thể hiện 3.500.000 VND."
    cases.append(tc07)

    tc08 = base_case("TC08", "Báo cáo và hóa đơn lệch số tiền")
    tc08.report_lines = [
        {"description": tc08.purpose, "invoice_ref": tc08.invoice_number, "amount": Decimal("3500000"), "note": "Đính kèm hóa đơn điện tử"}
    ]
    tc08.report_total = Decimal("3500000")
    tc08.advance_amount = Decimal("4000000")
    tc08.expected_decision = "REQUEST_INFO"
    tc08.expected_uncertainty = "FACTUAL_UNKNOWN"
    tc08.expected_primary_rule = "SET_PAPER_AMOUNT_001"
    tc08.expected_findings = ["SET_PAPER_AMOUNT_001"]
    tc08.rationale = "Báo cáo ghi 3.500.000 VND, hóa đơn ghi 3.240.000 VND."
    cases.append(tc08)

    tc09 = base_case("TC09", "Mã số thuế người mua thuộc công ty khác")
    tc09.buyer_tax_code = "7777777777"
    tc09.expected_decision = "ESCALATE"
    tc09.expected_uncertainty = "OUTSIDE_POLICY"
    tc09.expected_primary_rule = "SET_BUYER_TAX_001"
    tc09.expected_findings = ["SET_BUYER_TAX_001"]
    tc09.implementation_status = "GAP_EXPECTED"
    tc09.rationale = f"MST người mua là 7777777777, khác MST đã cấu hình của công ty {COMPANY_TAX_CODE}."
    cases.append(tc09)

    tc10 = base_case("TC10", "Hóa đơn thiếu mã số thuế người mua")
    tc10.buyer_tax_code = None
    tc10.expected_decision = "REQUEST_INFO"
    tc10.expected_uncertainty = "FACTUAL_UNKNOWN"
    tc10.expected_primary_rule = "SET_REQUIRED_FIELD_001"
    tc10.expected_findings = ["SET_REQUIRED_FIELD_001"]
    tc10.implementation_status = "GAP_EXPECTED"
    tc10.rationale = "Hóa đơn điện tử không có MST người mua."
    cases.append(tc10)

    tc11 = base_case("TC11", "Số học trên hóa đơn không nhất quán", total_kind="office")
    tc11.invoice_total = Decimal("5500000")
    tc11.report_total = Decimal("5500000")
    tc11.report_lines = [
        {"description": tc11.purpose, "invoice_ref": tc11.invoice_number, "amount": Decimal("5500000"), "note": "Đính kèm hóa đơn điện tử"}
    ]
    tc11.advance_amount = Decimal("6000000")
    tc11.expected_decision = "REQUEST_INFO"
    tc11.expected_uncertainty = "FACTUAL_UNKNOWN"
    tc11.expected_primary_rule = "SET_ARITH_001"
    tc11.expected_findings = ["SET_ARITH_001"]
    tc11.implementation_status = "GAP_EXPECTED"
    tc11.rationale = "Tiền hàng 5.000.000 + thuế 400.000 nhưng tổng hóa đơn ghi 5.500.000 VND."
    cases.append(tc11)

    tc12 = deepcopy(tc01)
    tc12.case_id = "TC12"
    tc12.title = "Hóa đơn trùng định danh với hồ sơ trước"
    tc12.purpose = "Chi phí tiếp khách lần nộp thứ hai"
    tc12.expected_decision = "ESCALATE"
    tc12.expected_uncertainty = "SUSPICIOUS"
    tc12.expected_primary_rule = "SET_DUP_001"
    tc12.expected_findings = ["SET_DUP_001"]
    tc12.implementation_status = "GAP_EXPECTED"
    tc12.rationale = "Giữ nguyên số hóa đơn, người bán, ngày và tổng tiền của TC01 để kiểm tra semantic duplicate."
    cases.append(tc12)

    tc13 = TestCase(
        case_id="TC13",
        title="Khoản chi rượu bia cần chuyển policy",
        purpose="Mua đồ uống phục vụ cá nhân sau giờ làm",
        invoice_number="00001301",
        lines=[Line("Rượu vang nhập khẩu", "Chai", Decimal("2"), Decimal("1000000"))],
        tax_rate=Decimal("8"),
        advance_amount=Decimal("3000000"),
        expected_decision="ESCALATE",
        expected_uncertainty="OUTSIDE_POLICY",
        expected_primary_rule="SET_CATEGORY_001",
        expected_findings=["SET_CATEGORY_001"],
        implementation_status="GAP_EXPECTED",
        rationale="Nội dung thể hiện chi tiêu cá nhân và mặt hàng rượu.",
    )
    cases.append(tc13)

    tc14 = TestCase(
        case_id="TC14",
        title="Tổng hồ sơ vượt thẩm quyền 100 triệu",
        purpose="Mua máy chủ phục vụ hệ thống nội bộ",
        invoice_number="00001401",
        lines=[Line("Máy chủ doanh nghiệp", "Bộ", Decimal("1"), Decimal("110000000"))],
        tax_rate=Decimal("10"),
        advance_amount=Decimal("120000000"),
        expected_decision="ESCALATE",
        expected_uncertainty="BEYOND_AUTHORITY",
        expected_primary_rule="SET_AUTH_001",
        expected_findings=["SET_AUTH_001"],
        rationale="Tổng thanh toán 121.000.000 VND vượt hạn mức phê duyệt 100.000.000 VND.",
    )
    cases.append(tc14)

    tc15 = TestCase(
        case_id="TC15",
        title="Evidence cùng số tiền nhưng sai nội dung",
        purpose="Sửa chữa máy in văn phòng",
        invoice_number="00001501",
        lines=[Line("Vé máy bay khứ hồi", "Vé", Decimal("1"), Decimal("2500000"))],
        tax_rate=Decimal("8"),
        advance_amount=Decimal("3000000"),
        report_lines=[
            {"description": "Sửa chữa máy in văn phòng", "invoice_ref": "SC-2026-15", "amount": Decimal("2700000"), "note": "Đính kèm chứng từ"}
        ],
        report_total=Decimal("2700000"),
        expected_decision="REQUEST_INFO",
        expected_uncertainty="FACTUAL_UNKNOWN",
        expected_primary_rule="SET_PAPER_MATCH_001",
        expected_findings=["SET_PAPER_MATCH_001"],
        rationale="Số tiền trùng nhưng nội dung báo cáo là sửa máy in, hóa đơn là vé máy bay và số chứng từ khác.",
    )
    cases.append(tc15)

    return cases


def _invoice_evidence(
    evidence_id: str,
    file_stem: str,
    reference: str,
    description: str,
    subtotal: Decimal | str,
    *,
    tax_rate: Decimal | str = Decimal("8"),
    lines: list[Line] | None = None,
) -> EvidenceDocument:
    vendor_name, vendor_tax_code, vendor_address, vendor_signer = VENDOR_PROFILES.get(
        file_stem,
        (SELLER_NAME, SELLER_TAX_CODE, SELLER_ADDRESS, SELLER_SIGNER_NAME),
    )
    invoice_lines = lines or [
        Line(description, "Lần", Decimal("1"), Decimal(str(subtotal)))
    ]
    return EvidenceDocument(
        evidence_id=evidence_id,
        file_stem=file_stem,
        document_type="E_INVOICE",
        reference=reference,
        title="Hóa đơn giá trị gia tăng điện tử",
        description=description,
        lines=invoice_lines,
        tax_rate=Decimal(str(tax_rate)),
        issuer=vendor_name,
        seller_tax_code=vendor_tax_code,
        issuer_address=vendor_address,
        issuer_signer=vendor_signer,
    )


def _alternative_evidence(
    evidence_id: str,
    file_stem: str,
    document_type: str,
    reference: str,
    title: str,
    description: str,
    amount: Decimal | str,
    *,
    issuer: str,
) -> EvidenceDocument:
    return EvidenceDocument(
        evidence_id=evidence_id,
        file_stem=file_stem,
        document_type=document_type,
        reference=reference,
        title=title,
        description=description,
        amount=Decimal(str(amount)),
        issuer=issuer,
    )


def _report_line(
    line_id: str,
    description: str,
    reference: str,
    amount: Decimal | str,
    note: str,
) -> dict[str, Any]:
    return {
        "line_id": line_id,
        "description": description,
        "invoice_ref": reference,
        "amount": Decimal(str(amount)),
        "note": note,
    }


def _line_expectation(
    line_id: str,
    support_type: str,
    document_files: list[str],
    match_status: str,
    reason: str,
) -> dict[str, Any]:
    return {
        "report_line_id": line_id,
        "support_type": support_type,
        "document_files": document_files,
        "match_status": match_status,
        "reason": reason,
    }


def build_multi_document_cases() -> list[MultiDocumentCase]:
    cases: list[MultiDocumentCase] = []

    tc16_evidence = [
        _invoice_evidence("EV01", "hoa_don_khach_san", "00001601", "Phòng nghỉ công tác", "4000000"),
        _invoice_evidence("EV02", "hoa_don_ve_may_bay", "00001602", "Vé máy bay khứ hồi", "5000000"),
        _invoice_evidence("EV03", "hoa_don_taxi", "00001603", "Dịch vụ taxi sân bay", "500000"),
    ]
    cases.append(MultiDocumentCase(
        case_id="TC16",
        title="Hợp lệ ba hóa đơn trong một báo cáo công tác",
        purpose="Hoàn ứng chi phí chuyến công tác Đà Nẵng",
        report_lines=[
            _report_line("L01", "Phòng nghỉ công tác", "00001601", "4320000", "Đính kèm hóa đơn"),
            _report_line("L02", "Vé máy bay khứ hồi", "00001602", "5400000", "Đính kèm hóa đơn"),
            _report_line("L03", "Taxi sân bay", "00001603", "540000", "Đính kèm hóa đơn"),
        ],
        evidence_documents=tc16_evidence,
        advance_amount=Decimal("12000000"),
        expected_findings=["SET_MULTI_EVIDENCE_MATCH_001", "SET_PAPER_TOTAL_001", "SET_SETTLEMENT_DIFF_001"],
        rationale="Ba dòng được ghép một-một với ba hóa đơn; tổng hồ sơ 10.260.000 VND.",
        line_expectations=[
            _line_expectation("L01", "E_INVOICE", ["hoa_don_khach_san.pdf"], "MATCHED", "Số hóa đơn, nội dung và số tiền khớp."),
            _line_expectation("L02", "E_INVOICE", ["hoa_don_ve_may_bay.pdf"], "MATCHED", "Số hóa đơn, nội dung và số tiền khớp."),
            _line_expectation("L03", "E_INVOICE", ["hoa_don_taxi.pdf"], "MATCHED", "Số hóa đơn, nội dung và số tiền khớp."),
        ],
    ))

    tc17_evidence = [
        _invoice_evidence("EV01", "hoa_don_nha_hang", "00001701", "Dịch vụ ăn uống tiếp khách", "3000000"),
        _invoice_evidence("EV02", "hoa_don_phong_hop", "00001702", "Thuê phòng họp", "2000000"),
        _alternative_evidence("EV03", "ve_dien_tu_bai_do_xe", "E_TICKET", "VDT-171026-0385", "Vé điện tử bãi đỗ xe", "Phí gửi xe phục vụ buổi tiếp khách", "200000", issuer="BÃI ĐỖ XE THÀNH CÔNG"),
    ]
    cases.append(MultiDocumentCase(
        case_id="TC17",
        title="Hợp lệ hóa đơn kèm vé điện tử thay thế",
        purpose="Hoàn ứng chi phí tiếp khách và tổ chức cuộc họp",
        report_lines=[
            _report_line("L01", "Ăn uống tiếp khách", "00001701", "3240000", "Đính kèm hóa đơn"),
            _report_line("L02", "Thuê phòng họp", "00001702", "2160000", "Đính kèm hóa đơn"),
            _report_line("L03", "Phí gửi xe", "VDT-171026-0385", "200000", "Vé điện tử thay thế hóa đơn"),
        ],
        evidence_documents=tc17_evidence,
        advance_amount=Decimal("6000000"),
        expected_findings=["SET_MULTI_EVIDENCE_MATCH_001", "SET_ALTERNATIVE_DOCUMENT_001"],
        rationale="Hai hóa đơn và một vé điện tử hỗ trợ đủ ba dòng; tổng 5.600.000 VND.",
        line_expectations=[
            _line_expectation("L01", "E_INVOICE", ["hoa_don_nha_hang.pdf"], "MATCHED", "Hóa đơn khớp dòng tiếp khách."),
            _line_expectation("L02", "E_INVOICE", ["hoa_don_phong_hop.pdf"], "MATCHED", "Hóa đơn khớp dòng thuê phòng."),
            _line_expectation("L03", "ALTERNATIVE_DOCUMENT", ["ve_dien_tu_bai_do_xe.pdf"], "MATCHED", "Vé điện tử là tài liệu thay thế và khớp số tiền."),
        ],
    ))

    tc18_evidence = [
        _invoice_evidence("EV01", "hoa_don_linh_kien", "00001801", "Linh kiện sửa máy in", "2500000"),
        _alternative_evidence("EV02", "bien_ban_nghiem_thu", "ACCEPTANCE_MINUTES", "BBNT-181026-02", "Biên bản nghiệm thu dịch vụ", "Nhân công sửa chữa máy in", "1500000", issuer="TRUNG TÂM KỸ THUẬT PHÚC AN"),
        _alternative_evidence("EV03", "xac_nhan_chuyen_khoan", "BANK_TRANSFER", "FT262918104578", "Xác nhận giao dịch ngân hàng", "Thanh toán nhân công sửa chữa máy in", "1500000", issuer="NGÂN HÀNG TMCP NGOẠI THƯƠNG VIỆT NAM"),
    ]
    cases.append(MultiDocumentCase(
        case_id="TC18",
        title="Hợp lệ một dòng dùng hai tài liệu thay thế",
        purpose="Hoàn ứng sửa chữa máy in văn phòng",
        report_lines=[
            _report_line("L01", "Mua linh kiện máy in", "00001801", "2700000", "Đính kèm hóa đơn"),
            _report_line("L02", "Nhân công sửa chữa", "BBNT-181026-02 / FT262918104578", "1500000", "Biên bản nghiệm thu và chuyển khoản"),
        ],
        evidence_documents=tc18_evidence,
        advance_amount=Decimal("5000000"),
        expected_findings=["SET_MULTI_EVIDENCE_MATCH_001", "SET_ALTERNATIVE_DOCUMENT_001"],
        rationale="Dòng nhân công được hỗ trợ đồng thời bởi biên bản nghiệm thu và xác nhận chuyển khoản.",
        line_expectations=[
            _line_expectation("L01", "E_INVOICE", ["hoa_don_linh_kien.pdf"], "MATCHED", "Hóa đơn linh kiện khớp 2.700.000 VND."),
            _line_expectation("L02", "ALTERNATIVE_DOCUMENT", ["bien_ban_nghiem_thu.pdf", "xac_nhan_chuyen_khoan.pdf"], "MATCHED", "Hai tài liệu cùng xác nhận dịch vụ và khoản thanh toán 1.500.000 VND."),
        ],
    ))

    tc19_evidence = [
        _invoice_evidence("EV01", "hoa_don_khach_san", "00001901", "Phòng nghỉ công tác", "2000000"),
        _invoice_evidence("EV02", "hoa_don_taxi", "00001902", "Taxi nội thành", "500000"),
    ]
    cases.append(MultiDocumentCase(
        case_id="TC19",
        title="Thiếu chứng từ và không có lý do cho một dòng",
        purpose="Hoàn ứng chuyến khảo sát khách hàng",
        report_lines=[
            _report_line("L01", "Phòng nghỉ", "00001901", "2160000", "Đính kèm hóa đơn"),
            _report_line("L02", "Taxi nội thành", "00001902", "540000", "Đính kèm hóa đơn"),
            _report_line("L03", "Ăn trưa khi khảo sát", "Không có", "800000", ""),
        ],
        evidence_documents=tc19_evidence,
        advance_amount=Decimal("4000000"),
        expected_decision="REQUEST_INFO",
        expected_uncertainty="FACTUAL_UNKNOWN",
        expected_primary_rule="SET_EVIDENCE_REASON_001",
        expected_findings=["SET_EVIDENCE_REASON_001"],
        rationale="Dòng ăn trưa 800.000 VND không có chứng từ và không ghi lý do.",
        line_expectations=[
            _line_expectation("L01", "E_INVOICE", ["hoa_don_khach_san.pdf"], "MATCHED", "Hóa đơn khớp."),
            _line_expectation("L02", "E_INVOICE", ["hoa_don_taxi.pdf"], "MATCHED", "Hóa đơn khớp."),
            _line_expectation("L03", "NO_DOCUMENT", [], "MISSING_REASON", "Không có file và cột ghi chú không nêu lý do."),
        ],
    ))

    tc20_evidence = [
        _invoice_evidence("EV01", "hoa_don_nhien_lieu", "00002001", "Nhiên liệu đi khảo sát", "1000000"),
        _alternative_evidence("EV02", "ve_dien_tu_cao_toc", "E_TICKET", "ETC-201026-7742", "Vé điện tử cao tốc", "Phí đường bộ chuyến khảo sát", "300000", issuer="TRẠM THU PHÍ LONG THÀNH"),
    ]
    cases.append(MultiDocumentCase(
        case_id="TC20",
        title="Có lý do cho dòng không có chứng từ",
        purpose="Hoàn ứng chuyến khảo sát hiện trường",
        report_lines=[
            _report_line("L01", "Nhiên liệu", "00002001", "1080000", "Đính kèm hóa đơn"),
            _report_line("L02", "Phí cao tốc", "ETC-201026-7742", "300000", "Vé điện tử thay thế"),
            _report_line("L03", "Phí gửi xe hiện trường", "Không có", "100000", "Bãi giữ xe tạm không phát hành chứng từ"),
        ],
        evidence_documents=tc20_evidence,
        advance_amount=Decimal("2000000"),
        expected_decision="ESCALATE",
        expected_uncertainty="OUTSIDE_POLICY",
        expected_primary_rule="SET_NO_INVOICE_EXCEPTION_001",
        expected_findings=["SET_ALTERNATIVE_DOCUMENT_001", "SET_NO_INVOICE_EXCEPTION_001"],
        rationale="Dòng không chứng từ đã có lý do nhưng vẫn cần kế toán duyệt ngoại lệ, không được tự động bác bỏ.",
        line_expectations=[
            _line_expectation("L01", "E_INVOICE", ["hoa_don_nhien_lieu.pdf"], "MATCHED", "Hóa đơn khớp."),
            _line_expectation("L02", "ALTERNATIVE_DOCUMENT", ["ve_dien_tu_cao_toc.pdf"], "MATCHED", "Vé điện tử khớp."),
            _line_expectation("L03", "NO_DOCUMENT", [], "MANUAL_EXCEPTION", "Có lý do cụ thể; kế toán quyết định chấp nhận ngoại lệ."),
        ],
    ))

    tc21_evidence = [
        _invoice_evidence("EV01", "hoa_don_thue_hoi_truong", "00002101", "Thuê hội trường", "10000000"),
        _invoice_evidence("EV02", "hoa_don_in_an", "00002102", "In tài liệu hội nghị", "2000000"),
    ]
    cases.append(MultiDocumentCase(
        case_id="TC21",
        title="Một hóa đơn được khai báo nhưng chưa đính kèm",
        purpose="Hoàn ứng chi phí tổ chức hội nghị",
        report_lines=[
            _report_line("L01", "Thuê hội trường", "00002101", "10800000", "Đính kèm hóa đơn"),
            _report_line("L02", "In tài liệu", "00002102", "2160000", "Đính kèm hóa đơn"),
            _report_line("L03", "Nước uống hội nghị", "00002103", "1080000", "Đính kèm hóa đơn"),
        ],
        evidence_documents=tc21_evidence,
        advance_amount=Decimal("15000000"),
        expected_decision="REQUEST_INFO",
        expected_uncertainty="FACTUAL_UNKNOWN",
        expected_primary_rule="SET_EVIDENCE_MISSING_001",
        expected_findings=["SET_EVIDENCE_MISSING_001"],
        rationale="Report tham chiếu hóa đơn 00002103 nhưng hồ sơ không có file tương ứng.",
        line_expectations=[
            _line_expectation("L01", "E_INVOICE", ["hoa_don_thue_hoi_truong.pdf"], "MATCHED", "Hóa đơn khớp."),
            _line_expectation("L02", "E_INVOICE", ["hoa_don_in_an.pdf"], "MATCHED", "Hóa đơn khớp."),
            _line_expectation("L03", "E_INVOICE", [], "MISSING_DOCUMENT", "Có số hóa đơn 00002103 nhưng thiếu file tương ứng."),
        ],
    ))

    tc22_evidence = [
        _invoice_evidence("EV01", "hoa_don_thiet_bi", "00002201", "Thiết bị mạng văn phòng", "5000000"),
        _invoice_evidence("EV02", "hoa_don_van_chuyen", "00002202", "Vận chuyển thiết bị", "500000"),
    ]
    cases.append(MultiDocumentCase(
        case_id="TC22",
        title="Một hóa đơn bị dùng vượt giá trị cho hai dòng",
        purpose="Hoàn ứng mua và lắp đặt thiết bị mạng",
        report_lines=[
            _report_line("L01", "Mua thiết bị mạng", "00002201", "5400000", "Đính kèm hóa đơn"),
            _report_line("L02", "Nhân công lắp đặt", "00002201", "2160000", "Dùng chung hóa đơn 00002201"),
            _report_line("L03", "Vận chuyển", "00002202", "540000", "Đính kèm hóa đơn"),
        ],
        evidence_documents=tc22_evidence,
        advance_amount=Decimal("9000000"),
        expected_decision="REQUEST_INFO",
        expected_uncertainty="FACTUAL_UNKNOWN",
        expected_primary_rule="SET_EVIDENCE_REUSE_001",
        expected_findings=["SET_EVIDENCE_REUSE_001", "SET_PAPER_AMOUNT_001"],
        rationale="Hóa đơn 00002201 có tổng 5.400.000 VND nhưng hai dòng cùng tham chiếu với tổng 7.560.000 VND.",
        line_expectations=[
            _line_expectation("L01", "E_INVOICE", ["hoa_don_thiet_bi.pdf"], "CONFLICT", "Riêng dòng này bằng toàn bộ giá trị hóa đơn 00002201."),
            _line_expectation("L02", "E_INVOICE", ["hoa_don_thiet_bi.pdf"], "CONFLICT", "Tái sử dụng hóa đơn 00002201 làm tổng phân bổ vượt giá trị hóa đơn."),
            _line_expectation("L03", "E_INVOICE", ["hoa_don_van_chuyen.pdf"], "MATCHED", "Hóa đơn vận chuyển khớp."),
        ],
    ))

    stationery_lines = [
        Line("Giấy in A4", "Thùng", Decimal("2"), Decimal("1000000")),
        Line("Bút và dụng cụ văn phòng", "Lô", Decimal("1"), Decimal("1000000")),
    ]
    tc23_evidence = [
        _invoice_evidence("EV01", "hoa_don_van_phong_pham", "00002301", "Văn phòng phẩm", "3000000", lines=stationery_lines),
        _invoice_evidence("EV02", "hoa_don_muc_in", "00002302", "Mực in văn phòng", "5000000"),
    ]
    cases.append(MultiDocumentCase(
        case_id="TC23",
        title="Hợp lệ một hóa đơn phân bổ cho nhiều dòng",
        purpose="Hoàn ứng mua vật tư văn phòng",
        report_lines=[
            _report_line("L01", "Giấy in A4", "00002301", "2160000", "Một phần hóa đơn 00002301"),
            _report_line("L02", "Bút và dụng cụ", "00002301", "1080000", "Một phần hóa đơn 00002301"),
            _report_line("L03", "Mực in", "00002302", "5400000", "Đính kèm hóa đơn"),
        ],
        evidence_documents=tc23_evidence,
        advance_amount=Decimal("9000000"),
        expected_findings=["SET_MULTI_EVIDENCE_MATCH_001", "SET_EVIDENCE_ALLOCATION_001"],
        rationale="Hai dòng phân bổ hóa đơn 00002301 cộng đúng 3.240.000 VND; hóa đơn 00002302 khớp riêng 5.400.000 VND.",
        line_expectations=[
            _line_expectation("L01", "E_INVOICE", ["hoa_don_van_phong_pham.pdf"], "MATCHED_ALLOCATION", "Cùng L02 phân bổ đủ giá trị hóa đơn 00002301."),
            _line_expectation("L02", "E_INVOICE", ["hoa_don_van_phong_pham.pdf"], "MATCHED_ALLOCATION", "Cùng L01 phân bổ đủ giá trị hóa đơn 00002301."),
            _line_expectation("L03", "E_INVOICE", ["hoa_don_muc_in.pdf"], "MATCHED", "Hóa đơn khớp."),
        ],
    ))

    tc24_evidence = [
        _invoice_evidence("EV01", "hoa_don_khach_san", "00002401", "Phòng nghỉ công tác", "2000000"),
        _invoice_evidence("EV02", "hoa_don_taxi", "00002402", "Taxi công tác", "500000"),
        _invoice_evidence("EV03", "hoa_don_khong_khai_bao", "00002409", "Ăn uống cá nhân", "1000000"),
    ]
    cases.append(MultiDocumentCase(
        case_id="TC24",
        title="Có hóa đơn đính kèm nhưng không xuất hiện trên report",
        purpose="Hoàn ứng chuyến công tác ngắn ngày",
        report_lines=[
            _report_line("L01", "Phòng nghỉ", "00002401", "2160000", "Đính kèm hóa đơn"),
            _report_line("L02", "Taxi công tác", "00002402", "540000", "Đính kèm hóa đơn"),
        ],
        evidence_documents=tc24_evidence,
        advance_amount=Decimal("3000000"),
        expected_decision="REQUEST_INFO",
        expected_uncertainty="FACTUAL_UNKNOWN",
        expected_primary_rule="SET_UNREFERENCED_EVIDENCE_001",
        expected_findings=["SET_UNREFERENCED_EVIDENCE_001"],
        rationale="Hóa đơn 00002409 được tải lên nhưng không có dòng report nào tham chiếu; cần xác nhận bỏ nhầm hay đính kèm nhầm.",
        line_expectations=[
            _line_expectation("L01", "E_INVOICE", ["hoa_don_khach_san.pdf"], "MATCHED", "Hóa đơn khớp."),
            _line_expectation("L02", "E_INVOICE", ["hoa_don_taxi.pdf"], "MATCHED", "Hóa đơn khớp."),
        ],
    ))

    tc25_evidence = [
        _invoice_evidence("EV01", "hoa_don_ve_tau", "00002501", "Vé tàu công tác", "1000000"),
        _invoice_evidence("EV02", "hoa_don_khach_san", "00002502", "Phòng nghỉ công tác", "2000000"),
        _alternative_evidence("EV03", "phieu_thu_di_chuyen", "RECEIPT", "PT-251026-016", "Phiếu thu dịch vụ di chuyển", "Di chuyển nội thành", "700000", issuer="HỢP TÁC XÃ VẬN TẢI THÀNH ĐẠT"),
    ]
    cases.append(MultiDocumentCase(
        case_id="TC25",
        title="Tài liệu thay thế lệch số tiền với report",
        purpose="Hoàn ứng chi phí công tác tỉnh",
        report_lines=[
            _report_line("L01", "Vé tàu", "00002501", "1080000", "Đính kèm hóa đơn"),
            _report_line("L02", "Phòng nghỉ", "00002502", "2160000", "Đính kèm hóa đơn"),
            _report_line("L03", "Di chuyển nội thành", "PT-251026-016", "900000", "Phiếu thu thay thế hóa đơn"),
        ],
        evidence_documents=tc25_evidence,
        advance_amount=Decimal("5000000"),
        expected_decision="REQUEST_INFO",
        expected_uncertainty="FACTUAL_UNKNOWN",
        expected_primary_rule="SET_ALTERNATIVE_AMOUNT_001",
        expected_findings=["SET_ALTERNATIVE_AMOUNT_001"],
        rationale="Report ghi 900.000 VND nhưng phiếu thu PT-251026-016 chỉ ghi 700.000 VND.",
        line_expectations=[
            _line_expectation("L01", "E_INVOICE", ["hoa_don_ve_tau.pdf"], "MATCHED", "Hóa đơn khớp."),
            _line_expectation("L02", "E_INVOICE", ["hoa_don_khach_san.pdf"], "MATCHED", "Hóa đơn khớp."),
            _line_expectation("L03", "ALTERNATIVE_DOCUMENT", ["phieu_thu_di_chuyen.pdf"], "AMOUNT_CONFLICT", "Lệch 200.000 VND giữa report và phiếu thu."),
        ],
    ))

    return cases


def expected_payload(case: TestCase) -> dict[str, Any]:
    report_sum = sum(
        (Decimal(str(line["amount"])) for line in case.settlement_lines()),
        Decimal("0"),
    )
    return {
        "schema_version": "1.0",
        "case_id": case.case_id,
        "title": case.title,
        "synthetic_test_data": True,
        "submission": {
            "settlement_type": "ADVANCE_SETTLEMENT",
            "source_type": "PAPER_SCAN",
            "request_document": "report_hoan_ung.pdf",
            "evidence_documents": ["hoa_don_dien_tu.pdf"],
        },
        "expected": {
            "automation_decision": case.expected_decision,
            "uncertainty_type": case.expected_uncertainty,
            "primary_rule": case.expected_primary_rule,
            "findings": case.expected_findings,
            "followup": (
                "PENDING_APPROVAL"
                if case.expected_decision == "AUTO_PROCESS"
                else None
            ),
        },
        "implementation_status": case.implementation_status,
        "accounting_facts": {
            "report_line_sum": str(report_sum),
            "report_total": str(case.settlement_total()),
            "advance_amount": str(case.advance_amount),
            "declared_difference": str(case.difference()),
            "invoice_subtotal": str(case.subtotal()),
            "invoice_tax": str(case.tax()),
            "invoice_total": str(case.total()),
            "buyer_tax_code": case.buyer_tax_code,
            "expected_company_tax_code": COMPANY_TAX_CODE,
        },
        "rationale": case.rationale,
    }


def expected_multi_payload(case: MultiDocumentCase) -> dict[str, Any]:
    report_sum = sum(
        (Decimal(str(line["amount"])) for line in case.report_lines),
        Decimal("0"),
    )
    evidence_files = [f"{document.file_stem}.pdf" for document in case.evidence_documents]
    evidence_facts = [
        {
            "evidence_id": document.evidence_id,
            "file": f"{document.file_stem}.pdf",
            "document_type": document.document_type,
            "reference": document.reference,
            "description": document.description,
            "amount": str(document.total()),
        }
        for document in case.evidence_documents
    ]
    report_lines = [
        {
            **line,
            "amount": str(line["amount"]),
        }
        for line in case.report_lines
    ]
    return {
        "schema_version": "2.0",
        "case_id": case.case_id,
        "title": case.title,
        "synthetic_test_data": True,
        "submission": {
            "settlement_type": "ADVANCE_SETTLEMENT",
            "source_type": "PAPER_SCAN",
            "request_document": "report_hoan_ung.pdf",
            "evidence_documents": evidence_files,
        },
        "expected": {
            "automation_decision": case.expected_decision,
            "uncertainty_type": case.expected_uncertainty,
            "primary_rule": case.expected_primary_rule,
            "findings": case.expected_findings,
            "followup": (
                "PENDING_APPROVAL"
                if case.expected_decision == "AUTO_PROCESS"
                else None
            ),
        },
        "implementation_status": case.implementation_status,
        "accounting_facts": {
            "report_lines": report_lines,
            "report_line_sum": str(report_sum),
            "report_total": str(case.settlement_total()),
            "advance_amount": str(case.advance_amount),
            "declared_difference": str(case.difference()),
            "evidence": evidence_facts,
        },
        "line_support_expectations": case.line_expectations,
        "rationale": case.rationale,
    }


def write_readme(
    root: Path,
    cases: list[TestCase | MultiDocumentCase],
) -> None:
    lines = [
        "# Bộ testcase UC-03 hoàn ứng và quyết toán tạm ứng",
        "",
        "Toàn bộ tài liệu là dữ liệu giả lập, không có giá trị pháp lý hoặc kế toán.",
        "",
        "## Cách nộp một case",
        "",
        "- Chọn `Quyết toán tạm ứng` và `Đơn giấy`.",
        "- Tải `report_hoan_ung.pdf` làm đơn/báo cáo chính.",
        "- Tải toàn bộ file PDF còn lại trong thư mục case làm evidence.",
        "- `expected.json` là oracle nghiệp vụ, không tải vào hệ thống.",
        "- File DOCX là nguồn có thể chỉnh sửa; PDF là bản dùng để OCR.",
        "- Từ TC16, một report có nhiều dòng và nhiều evidence. Evidence có thể là hóa đơn,",
        "  tài liệu thay thế hoặc không tồn tại nếu report đã khai báo lý do.",
        "- `line_support_expectations` trong `expected.json` là mapping chuẩn giữa từng dòng",
        "  report và các file chứng minh tương ứng.",
        "",
        "## Danh sách case",
        "",
        "| Case | Tình huống | Kết quả mong đợi | Trạng thái triển khai |",
        "| --- | --- | --- | --- |",
    ]
    for case in cases:
        result = case.expected_decision
        if case.expected_primary_rule:
            result += f" / {case.expected_primary_rule}"
        lines.append(
            f"| {case.case_id} | {case.title} | {result} | {case.implementation_status} |"
        )
    lines.extend(
        [
            "",
            "`GAP_EXPECTED` là testcase chủ đích cho rule chưa triển khai đầy đủ. Case này",
            "được giữ trong bộ regression để tránh nhầm khoảng trống sản phẩm thành happy path.",
            "",
        ]
    )
    (root / "README.md").write_text("\n".join(lines), encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("data/testcase_uc3"),
    )
    parser.add_argument("--force", action="store_true")
    args = parser.parse_args()
    root = args.output.resolve()
    if root.exists():
        if not args.force:
            raise SystemExit(f"Output đã tồn tại: {root}. Dùng --force để tạo lại.")
        shutil.rmtree(root)
    root.mkdir(parents=True)
    _register_pdf_fonts()
    cases = build_cases()
    multi_cases = build_multi_document_cases()
    manifest: list[dict[str, Any]] = []
    for case in cases:
        case_dir = root / f"{case.case_id}_{case.title.lower().replace(' ', '_')}"
        case_dir.mkdir()
        create_report_docx(case, case_dir / "report_hoan_ung.docx")
        create_invoice_docx(case, case_dir / "hoa_don_dien_tu.docx")
        create_report_pdf(case, case_dir / "report_hoan_ung.pdf")
        create_invoice_pdf(case, case_dir / "hoa_don_dien_tu.pdf")
        payload = expected_payload(case)
        (case_dir / "expected.json").write_text(
            json.dumps(payload, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        manifest.append(
            {
                "case_id": case.case_id,
                "folder": case_dir.name,
                "title": case.title,
                "expected_decision": case.expected_decision,
                "expected_primary_rule": case.expected_primary_rule,
                "implementation_status": case.implementation_status,
            }
        )
    for case in multi_cases:
        case_dir = root / f"{case.case_id}_{case.title.lower().replace(' ', '_')}"
        case_dir.mkdir()
        report_case = case.as_report_case()
        create_report_docx(report_case, case_dir / "report_hoan_ung.docx")
        create_report_pdf(report_case, case_dir / "report_hoan_ung.pdf")
        for evidence in case.evidence_documents:
            docx_path = case_dir / f"{evidence.file_stem}.docx"
            pdf_path = case_dir / f"{evidence.file_stem}.pdf"
            if evidence.document_type == "E_INVOICE":
                invoice_case = evidence_as_invoice_case(case, evidence)
                create_invoice_docx(invoice_case, docx_path)
                create_invoice_pdf(invoice_case, pdf_path)
            else:
                create_alternative_docx(case, evidence, docx_path)
                create_alternative_pdf(case, evidence, pdf_path)
        payload = expected_multi_payload(case)
        (case_dir / "expected.json").write_text(
            json.dumps(payload, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        manifest.append(
            {
                "case_id": case.case_id,
                "folder": case_dir.name,
                "title": case.title,
                "expected_decision": case.expected_decision,
                "expected_primary_rule": case.expected_primary_rule,
                "implementation_status": case.implementation_status,
                "evidence_count": len(case.evidence_documents),
            }
        )
    (root / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    all_cases: list[TestCase | MultiDocumentCase] = [*cases, *multi_cases]
    write_readme(root, all_cases)
    print(f"Generated {len(all_cases)} test cases at {root}")


if __name__ == "__main__":
    main()
